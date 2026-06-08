# File này xử lý logic chính cho lịch sử hoạt động:
# - Tạo log hoạt động
# - Lọc danh sách hoạt động
# - Phân quyền xem log
# - Thống kê hoạt động
# - Xuất Excel
# - Lọc theo thời gian tạo hoạt động

from bson import ObjectId
from pymongo import DESCENDING
import unicodedata
import re
from io import BytesIO
from datetime import datetime, timedelta, time, timezone
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from mongo import activities_collection, users_collection

from templates.activity.activity_model import (
    activity_serializer,
    create_activity_model,
    is_valid_object_id,
)


# Múi giờ dùng cho giao diện Việt Nam
APP_TIMEZONE = ZoneInfo("Asia/Ho_Chi_Minh")


# Các role này được xem toàn bộ log hoạt động
ROLES_CAN_VIEW_ALL = ["ADMIN", "QUAN_LY"]

# Role nhân viên thường, chỉ xem log của chính mình
ROLE_EMPLOYEE = "NHAN_VIEN"

# Số log hiển thị mặc định trên mỗi trang
DEFAULT_ACTIVITY_LIMIT = 10

# Các loại hoạt động hiển thị trên bộ lọc
ACTIVITY_TYPE_OPTIONS = [
    "Báo cáo",
    "Cấp phát",
    "Thu hồi",
    "Tài sản",
    "Người dùng",
    "Phân quyền",
    "Mail",
    "Hoạt động",
    "Hệ thống",
]


# Các module/path không cần ghi lịch sử hoạt động
# Notification chỉ là thông báo nội bộ, ví dụ bấm Xem để đánh dấu đã đọc
# nên không cần sinh log "Cập nhật dữ liệu notifications"
SKIP_ACTIVITY_LOG_MODULES = [
    "notification",
    "notifications",
]

SKIP_ACTIVITY_LOG_PATHS = [
    "/api/notifications",
    "/notifications",
    "/socket.io",
]


# Kiểm tra có nên bỏ qua ghi log hay không
def should_skip_activity_log(module=None, path=None, action=None):
    module_text = str(module or "").strip().lower()
    path_text = str(path or "").strip().lower()
    action_text = str(action or "").strip().lower()

    if module_text in SKIP_ACTIVITY_LOG_MODULES:
        return True

    for skip_path in SKIP_ACTIVITY_LOG_PATHS:
        if path_text.startswith(skip_path):
            return True

    if "notification" in path_text or "notifications" in path_text:
        return True

    if "notification" in action_text or "notifications" in action_text:
        return True

    return False


# Tìm người dùng hiện tại theo user_id
def get_current_user_by_id(user_id):
    if not user_id:
        return None

    if not is_valid_object_id(user_id):
        return None

    return users_collection.find_one({"_id": ObjectId(user_id)})


# Kiểm tra role có quyền xem tất cả hoạt động hay không
def can_view_all_activities(role):
    return role in ROLES_CAN_VIEW_ALL


# Chuẩn hóa text: viết thường và bỏ dấu tiếng Việt để tìm kiếm dễ hơn
def normalize_text(value):
    if value is None:
        return ""

    text = str(value).strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")

    return text


# Lấy giá trị đầu tiên có tồn tại trong request args
# Dùng để hỗ trợ nhiều tên input khác nhau từ giao diện
def get_first_arg_value(args, *keys):
    for key in keys:
        value = args.get(key)

        if value not in [None, ""]:
            return value

    return ""


# Parse chuỗi ngày tháng từ giao diện thành datetime của Python
# Người dùng chọn ngày theo giờ Việt Nam
def parse_activity_datetime(value, is_end_of_day=False):
    if not value:
        return None

    if isinstance(value, datetime):
        return value

    text = str(value).strip()

    if not text or text == "Tất cả":
        return None

    clean_text = (
        text
        .replace("Z", "")
        .replace("+00:00", "")
        .strip()
    )

    date_only_formats = [
        "%Y-%m-%d",
        "%d/%m/%Y",
    ]

    datetime_formats = [
        "%Y-%m-%dT%H:%M:%S.%f",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
    ]

    for fmt in date_only_formats:
        try:
            parsed_date = datetime.strptime(clean_text, fmt).date()

            if is_end_of_day:
                return datetime.combine(parsed_date, time.max)

            return datetime.combine(parsed_date, time.min)
        except Exception:
            pass

    for fmt in datetime_formats:
        try:
            return datetime.strptime(clean_text[:26], fmt)
        except Exception:
            pass

    return None


# Chuyển datetime giờ Việt Nam sang UTC naive để query MongoDB
# MongoDB/PyMongo thường lưu datetime ở UTC dạng không có tzinfo
def convert_vietnam_time_to_utc_naive(value):
    if not value:
        return None

    if value.tzinfo is None:
        local_value = value.replace(tzinfo=APP_TIMEZONE)
    else:
        local_value = value.astimezone(APP_TIMEZONE)

    utc_value = local_value.astimezone(timezone.utc)

    return utc_value.replace(tzinfo=None)


# Chuyển datetime UTC từ database sang giờ Việt Nam để xuất Excel
def convert_utc_to_vietnam_time(value):
    if not value:
        return None

    if value.tzinfo is None:
        utc_value = value.replace(tzinfo=timezone.utc)
    else:
        utc_value = value.astimezone(timezone.utc)

    return utc_value.astimezone(APP_TIMEZONE)


# Tạo khoảng thời gian nhanh theo lựa chọn từ giao diện
def build_relative_time_range(time_filter):
    if not time_filter or time_filter == "Tất cả":
        return None, None

    normalized = normalize_text(time_filter)
    normalized = normalized.replace("_", " ").replace("-", " ")

    now = datetime.now(APP_TIMEZONE)
    today = now.date()

    today_start = datetime.combine(today, time.min)
    today_end = datetime.combine(today, time.max)

    if normalized in ["today", "hom nay"]:
        return today_start, today_end

    if normalized in ["yesterday", "hom qua"]:
        yesterday = today - timedelta(days=1)

        return (
            datetime.combine(yesterday, time.min),
            datetime.combine(yesterday, time.max),
        )

    if normalized in ["last 7 days", "7 days", "7 ngay", "7 ngay qua"]:
        start_date = today - timedelta(days=6)

        return (
            datetime.combine(start_date, time.min),
            today_end,
        )

    if normalized in ["last 30 days", "30 days", "30 ngay", "30 ngay qua"]:
        start_date = today - timedelta(days=29)

        return (
            datetime.combine(start_date, time.min),
            today_end,
        )

    if normalized in ["this week", "tuan nay"]:
        start_date = today - timedelta(days=now.weekday())

        return (
            datetime.combine(start_date, time.min),
            today_end,
        )

    if normalized in ["this month", "thang nay"]:
        start_date = today.replace(day=1)

        return (
            datetime.combine(start_date, time.min),
            today_end,
        )

    return None, None


# Tạo điều kiện lọc thời gian cho MongoDB
# Người dùng chọn ngày theo giờ Việt Nam
# Trước khi query MongoDB sẽ đổi sang UTC
def build_activity_time_condition(args):
    start_value = get_first_arg_value(
        args,
        "start_date",
        "date_from",
        "from_date",
        "tu_ngay",
    )

    end_value = get_first_arg_value(
        args,
        "end_date",
        "date_to",
        "to_date",
        "den_ngay",
    )

    time_filter = get_first_arg_value(
        args,
        "time_filter",
        "time_range",
        "date_range",
    )

    start_date_local = parse_activity_datetime(
        start_value,
        is_end_of_day=False,
    )

    end_date_local = parse_activity_datetime(
        end_value,
        is_end_of_day=True,
    )

    if not start_date_local and not end_date_local and time_filter:
        start_date_local, end_date_local = build_relative_time_range(time_filter)

    if start_date_local and end_date_local and start_date_local > end_date_local:
        start_date_local, end_date_local = end_date_local, start_date_local

    created_at_condition = {}

    if start_date_local:
        created_at_condition["$gte"] = convert_vietnam_time_to_utc_naive(
            start_date_local
        )

    if end_date_local:
        created_at_condition["$lte"] = convert_vietnam_time_to_utc_naive(
            end_date_local
        )

    if not created_at_condition:
        return {}

    return {
        "created_at": created_at_condition
    }


# Làm sạch metadata trước khi lưu log, ẩn các thông tin nhạy cảm như password/token
def clean_metadata(data):
    if not isinstance(data, dict):
        return {}

    hidden_fields = [
        "password",
        "password_hash",
        "confirm_password",
        "token",
        "access_token",
        "refresh_token",
    ]

    cleaned = {}

    for key, value in data.items():
        if key in hidden_fields:
            cleaned[key] = "***"
        else:
            cleaned[key] = value

    return cleaned


# Tự đoán module từ đường dẫn API, ví dụ /api/assets sẽ ra assets
def detect_module_from_path(path):
    if not path:
        return "system"

    parts = [part for part in path.split("/") if part]

    if len(parts) >= 2 and parts[0] == "api":
        return parts[1]

    if len(parts) >= 1:
        return parts[0]

    return "system"


# Tự tạo nội dung hành động dựa vào method và path
def build_action_from_request(method, path):
    module = detect_module_from_path(path)
    method = method.upper()

    module_name_map = {
        "assets": "tài sản",
        "assign": "cấp phát",
        "users": "người dùng",
        "mail": "mail",
        "permissions": "phân quyền",
        "reports": "báo cáo",
        "activities": "hoạt động",
        "system": "hệ thống",
        "notifications": "thông báo",
        "notification": "thông báo",
    }

    module_name = module_name_map.get(module, module)

    if method == "POST":
        return f"Tạo mới dữ liệu {module_name}"

    if method in ["PUT", "PATCH"]:
        return f"Cập nhật dữ liệu {module_name}"

    if method == "DELETE":
        return f"Xóa dữ liệu {module_name}"

    return f"Thao tác với dữ liệu {module_name}"


# Tạo 1 log hoạt động mới và lưu vào database
def create_activity_log(
    user_id,
    action,
    module=None,
    method=None,
    path=None,
    status_code=None,
    target_id=None,
    target_name=None,
    metadata=None,
):
    # Không ghi log cho notification
    # Ví dụ: bấm Xem notification sẽ cập nhật is_read/read_at,
    # thao tác này không cần lưu vào lịch sử hoạt động.
    if should_skip_activity_log(module=module, path=path, action=action):
        return {
            "success": True,
            "message": "Bỏ qua ghi log notification"
        }, 200

    current_user = get_current_user_by_id(user_id)

    if not current_user:
        return {
            "success": False,
            "message": "Không tìm thấy user tạo log"
        }, 404

    if not action:
        return {
            "success": False,
            "message": "Thiếu nội dung hoạt động"
        }, 400

    activity_data = {
        "user_id": current_user["_id"],
        "employee_code": current_user.get("employee_code"),
        "full_name": current_user.get("full_name"),
        "email": current_user.get("email"),
        "role": current_user.get("role"),

        "action": action,
        "module": module,
        "method": method,
        "path": path,
        "status_code": status_code,

        "target_id": target_id,
        "target_name": target_name,

        "metadata": clean_metadata(metadata or {}),
    }

    activity = create_activity_model(activity_data)

    result = activities_collection.insert_one(activity)

    created_activity = activities_collection.find_one({"_id": result.inserted_id})

    return {
        "success": True,
        "message": "Tạo log hoạt động thành công",
        "data": activity_serializer(created_activity)
    }, 201


# Tạo điều kiện tìm kiếm regex cho nhiều field và nhiều từ khóa
def make_regex_condition(fields, patterns):
    conditions = []

    for field in fields:
        for pattern in patterns:
            conditions.append({
                field: {
                    "$regex": pattern,
                    "$options": "i"
                }
            })

    if not conditions:
        return {}

    return {
        "$or": conditions
    }


# Tạo điều kiện lọc theo loại hoạt động như Báo cáo, Cấp phát, Thu hồi...
def build_activity_type_condition(activity_type):
    if not activity_type or activity_type == "Tất cả":
        return {}

    normalized = normalize_text(activity_type)

    action_fields = [
        "action",
        "target_name",
    ]

    general_fields = [
        "action",
        "module",
        "path",
        "target_name",
    ]

    if "bao cao" in normalized or "report" in normalized:
        return make_regex_condition(general_fields, [
            "báo cáo",
            "bao cao",
            "report",
            "reports",
        ])

    if "thu hoi" in normalized or "recover" in normalized or "recall" in normalized:
        return make_regex_condition(action_fields, [
            "thu hồi",
            "thu hoi",
            "recover",
            "recall",
            "return",
        ])

    if "cap phat" in normalized or "assign" in normalized:
        return {
            "$and": [
                make_regex_condition(action_fields, [
                    "cấp phát",
                    "cap phat",
                    "assign",
                    "assignment",
                ]),
                {
                    "action": {
                        "$not": {
                            "$regex": "thu hồi|thu hoi|recover|recall|return",
                            "$options": "i"
                        }
                    }
                },
                {
                    "target_name": {
                        "$not": {
                            "$regex": "thu hồi|thu hoi|recover|recall|return",
                            "$options": "i"
                        }
                    }
                }
            ]
        }

    if "tai san" in normalized or "asset" in normalized:
        return {
            "$and": [
                make_regex_condition(general_fields, [
                    "tài sản",
                    "tai san",
                    "asset",
                    "assets",
                ]),
                {
                    "action": {
                        "$not": {
                            "$regex": "cấp phát|cap phat|thu hồi|thu hoi|assign|recover|recall|return",
                            "$options": "i"
                        }
                    }
                }
            ]
        }

    if "nguoi dung" in normalized or "user" in normalized:
        return make_regex_condition(general_fields, [
            "người dùng",
            "nguoi dung",
            "user",
            "users",
        ])

    if "phan quyen" in normalized or "permission" in normalized:
        return make_regex_condition(general_fields, [
            "phân quyền",
            "phan quyen",
            "permission",
            "permissions",
        ])

    if "mail" in normalized or "email" in normalized:
        return make_regex_condition(general_fields, [
            "mail",
            "email",
        ])

    if "hoat dong" in normalized or "activity" in normalized:
        return make_regex_condition(general_fields, [
            "hoạt động",
            "hoat dong",
            "activity",
            "activities",
        ])

    if "he thong" in normalized or "system" in normalized:
        return make_regex_condition(general_fields, [
            "hệ thống",
            "he thong",
            "system",
        ])

    return make_regex_condition(general_fields, [activity_type])


# Gộp 2 query MongoDB lại với nhau bằng $and khi cần
def merge_query(base_query, extra_condition):
    if not extra_condition:
        return base_query

    if not base_query:
        return extra_condition

    return {
        "$and": [
            base_query,
            extra_condition,
        ]
    }


# Tạo query phân quyền: admin/quản lý xem được nhiều, nhân viên chỉ xem của mình
def build_permission_query(current_user, args):
    query = {}
    current_role = current_user.get("role")

    user_id = args.get("user_id")
    role = args.get("role")
    employee_code = args.get("employee_code")

    if can_view_all_activities(current_role):
        if user_id:
            if not is_valid_object_id(user_id):
                return None, {
                    "success": False,
                    "message": "user_id không hợp lệ"
                }, 400

            query["user_id"] = ObjectId(user_id)

        if role:
            query["role"] = role

        if employee_code:
            query["employee_code"] = employee_code

    else:
        query["user_id"] = current_user["_id"]

    return query, None, None


# Chuẩn hóa keyword tìm kiếm hoạt động
# Nếu người dùng nhập ACT-xxxxxx thì lấy phần xxxxxx để tìm theo đuôi ObjectId
def normalize_activity_keyword(value):
    text = str(value or "").strip()

    if not text:
        return ""

    if text.upper().startswith("ACT-"):
        return text[4:].strip()

    return text


# Tạo điều kiện regex trên field có thể là ObjectId bằng cách ép sang string
# Dùng để search được _id, user_id, target_id theo một phần id như ABC123
def make_to_string_regex_condition(field, pattern):
    return {
        "$expr": {
            "$regexMatch": {
                "input": {
                    "$toString": {
                        "$ifNull": [
                            f"${field}",
                            ""
                        ]
                    }
                },
                "regex": pattern,
                "options": "i"
            }
        }
    }


# Tạo query chính để lọc danh sách hoạt động từ request args
def build_activity_query(args, current_user):
    permission_query, error_response, error_status = build_permission_query(
        current_user,
        args,
    )

    if error_response:
        return None, error_response, error_status

    query = dict(permission_query)

    keyword = (
        args.get("keyword")
        or args.get("search")
        or args.get("q")
    )

    activity_type = (
        args.get("type")
        or args.get("activity_type")
        or args.get("loai")
    )

    user_name = (
        args.get("user_name")
        or args.get("full_name")
        or args.get("name")
        or args.get("ten_nguoi_dung")
    )

    module = args.get("module")
    action = args.get("action")
    status_code = args.get("status_code")

    extra_conditions = []

    if module and module != "Tất cả":
        query["module"] = {
            "$regex": module,
            "$options": "i"
        }

    if action and action != "Tất cả":
        query["action"] = {
            "$regex": action,
            "$options": "i"
        }

    if status_code and status_code != "Tất cả":
        try:
            query["status_code"] = int(status_code)
        except Exception:
            query["status_code"] = status_code

    if user_name and user_name != "Tất cả":
        query["full_name"] = {
            "$regex": user_name,
            "$options": "i"
        }

    type_condition = build_activity_type_condition(activity_type)

    if type_condition:
        extra_conditions.append(type_condition)

    time_condition = build_activity_time_condition(args)

    if time_condition:
        extra_conditions.append(time_condition)

    if keyword:
        keyword = normalize_activity_keyword(keyword)

        if keyword:
            keyword_pattern = re.escape(keyword)

            keyword_conditions = [
                {"employee_code": {"$regex": keyword_pattern, "$options": "i"}},
                {"full_name": {"$regex": keyword_pattern, "$options": "i"}},
                {"email": {"$regex": keyword_pattern, "$options": "i"}},
                {"role": {"$regex": keyword_pattern, "$options": "i"}},
                {"action": {"$regex": keyword_pattern, "$options": "i"}},
                {"module": {"$regex": keyword_pattern, "$options": "i"}},
                {"target_name": {"$regex": keyword_pattern, "$options": "i"}},
                {"path": {"$regex": keyword_pattern, "$options": "i"}},

                {"activity_code": {"$regex": keyword_pattern, "$options": "i"}},
                {"code": {"$regex": keyword_pattern, "$options": "i"}},
                {"log_code": {"$regex": keyword_pattern, "$options": "i"}},

                {"target_id": {"$regex": keyword_pattern, "$options": "i"}},
                {"metadata.target_id": {"$regex": keyword_pattern, "$options": "i"}},
                {"metadata.id": {"$regex": keyword_pattern, "$options": "i"}},

                make_to_string_regex_condition("_id", keyword_pattern),
                make_to_string_regex_condition("user_id", keyword_pattern),
                make_to_string_regex_condition("target_id", keyword_pattern),
                make_to_string_regex_condition("metadata.target_id", keyword_pattern),
                make_to_string_regex_condition("metadata.id", keyword_pattern),
            ]

            if is_valid_object_id(keyword):
                keyword_object_id = ObjectId(keyword)

                keyword_conditions.extend([
                    {"_id": keyword_object_id},
                    {"user_id": keyword_object_id},
                    {"target_id": keyword_object_id},
                    {"metadata.target_id": keyword_object_id},
                    {"metadata.id": keyword_object_id},
                ])

            extra_conditions.append({
                "$or": keyword_conditions
            })

    for condition in extra_conditions:
        query = merge_query(query, condition)

    return query, None, None


# Lấy danh sách hoạt động có phân trang và bộ lọc
def get_activities(args, current_user_id):
    current_user = get_current_user_by_id(current_user_id)

    if not current_user:
        return {
            "success": False,
            "message": "Không tìm thấy user hiện tại"
        }, 404

    try:
        page = int(args.get("page", 1))
    except Exception:
        page = 1

    if page < 1:
        page = 1

    limit = DEFAULT_ACTIVITY_LIMIT
    skip = (page - 1) * limit

    query, error_response, error_status = build_activity_query(args, current_user)

    if error_response:
        return error_response, error_status

    total = activities_collection.count_documents(query)

    activities = (
        activities_collection
        .find(query)
        .sort("created_at", DESCENDING)
        .skip(skip)
        .limit(limit)
    )

    data = [activity_serializer(activity) for activity in activities]

    total_pages = (total + limit - 1) // limit if total > 0 else 1

    current_role = current_user.get("role")

    return {
        "success": True,
        "message": "Lấy danh sách hoạt động thành công",
        "data": data,
        "viewer": {
            "id": str(current_user["_id"]),
            "employee_code": current_user.get("employee_code"),
            "full_name": current_user.get("full_name"),
            "email": current_user.get("email"),
            "role": current_role,
            "can_view_all": can_view_all_activities(current_role),
        },
        "filters": {
            "keyword": args.get("keyword") or args.get("search") or args.get("q") or "",
            "type": args.get("type") or args.get("activity_type") or args.get("loai") or "Tất cả",
            "user_name": args.get("user_name") or args.get("full_name") or args.get("name") or "",
            "module": args.get("module") or "Tất cả",
            "action": args.get("action") or "Tất cả",
            "status_code": args.get("status_code") or "Tất cả",
            "time_filter": args.get("time_filter") or args.get("time_range") or args.get("date_range") or "Tất cả",
            "start_date": args.get("start_date") or args.get("date_from") or args.get("from_date") or args.get("tu_ngay") or "",
            "end_date": args.get("end_date") or args.get("date_to") or args.get("to_date") or args.get("den_ngay") or "",
        },
        "pagination": {
            "page": page,
            "limit": limit,
            "total": total,
            "total_pages": total_pages,
            "has_prev": page > 1,
            "has_next": page < total_pages,
            "prev_page": page - 1 if page > 1 else 1,
            "next_page": page + 1 if page < total_pages else total_pages,
        }
    }, 200


# Lấy dữ liệu cho bộ lọc: loại hoạt động và danh sách người dùng
def get_activity_filter_options(current_user_id):
    current_user = get_current_user_by_id(current_user_id)

    if not current_user:
        return {
            "success": False,
            "message": "Không tìm thấy user hiện tại"
        }, 404

    permission_query, error_response, error_status = build_permission_query(
        current_user,
        {},
    )

    if error_response:
        return error_response, error_status

    type_options = []

    for activity_type in ACTIVITY_TYPE_OPTIONS:
        type_query = merge_query(
            dict(permission_query),
            build_activity_type_condition(activity_type)
        )

        count = activities_collection.count_documents(type_query)

        type_options.append({
            "label": activity_type,
            "value": activity_type,
            "count": count,
        })

    user_pipeline = [
        {
            "$match": permission_query
        },
        {
            "$group": {
                "_id": "$full_name",
                "count": {
                    "$sum": 1
                },
                "user_id": {
                    "$first": "$user_id"
                },
                "employee_code": {
                    "$first": "$employee_code"
                },
                "email": {
                    "$first": "$email"
                },
                "role": {
                    "$first": "$role"
                },
            }
        },
        {
            "$sort": {
                "_id": 1
            }
        }
    ]

    user_options = []

    for item in activities_collection.aggregate(user_pipeline):
        full_name = item.get("_id")

        if not full_name:
            continue

        user_id = item.get("user_id")

        user_options.append({
            "id": str(user_id) if user_id else "",
            "full_name": full_name,
            "value": full_name,
            "label": full_name,
            "employee_code": item.get("employee_code"),
            "email": item.get("email"),
            "role": item.get("role"),
            "count": item.get("count", 0),
        })

    total = activities_collection.count_documents(permission_query)

    return {
        "success": True,
        "message": "Lấy bộ lọc hoạt động thành công",
        "data": {
            "total": total,
            "type_options": type_options,
            "user_options": user_options,
            "limit": DEFAULT_ACTIVITY_LIMIT,
        }
    }, 200


# Lấy chi tiết 1 log hoạt động theo id
def get_activity_by_id(activity_id, current_user_id):
    current_user = get_current_user_by_id(current_user_id)

    if not current_user:
        return {
            "success": False,
            "message": "Không tìm thấy user hiện tại"
        }, 404

    if not is_valid_object_id(activity_id):
        return {
            "success": False,
            "message": "ID hoạt động không hợp lệ"
        }, 400

    activity = activities_collection.find_one({"_id": ObjectId(activity_id)})

    if not activity:
        return {
            "success": False,
            "message": "Không tìm thấy hoạt động"
        }, 404

    current_role = current_user.get("role")

    if not can_view_all_activities(current_role):
        if str(activity.get("user_id")) != str(current_user["_id"]):
            return {
                "success": False,
                "message": "Bạn không có quyền xem hoạt động này"
            }, 403

    return {
        "success": True,
        "message": "Lấy chi tiết hoạt động thành công",
        "data": activity_serializer(activity)
    }, 200


# Lấy thống kê số lượng log tạo mới, cập nhật, xóa
def get_activity_stats(current_user_id):
    current_user = get_current_user_by_id(current_user_id)

    if not current_user:
        return {
            "success": False,
            "message": "Không tìm thấy user hiện tại"
        }, 404

    current_role = current_user.get("role")

    query = {}

    if not can_view_all_activities(current_role):
        query["user_id"] = current_user["_id"]

    total = activities_collection.count_documents(query)

    create_count = activities_collection.count_documents({
        **query,
        "method": "POST"
    })

    update_count = activities_collection.count_documents({
        **query,
        "method": {"$in": ["PUT", "PATCH"]}
    })

    delete_count = activities_collection.count_documents({
        **query,
        "method": "DELETE"
    })

    return {
        "success": True,
        "message": "Lấy thống kê hoạt động thành công",
        "data": {
            "total": total,
            "create_count": create_count,
            "update_count": update_count,
            "delete_count": delete_count,
        }
    }, 200


# Xác định loại hoạt động khi xuất Excel
def detect_activity_export_type(activity):
    module = str(activity.get("module") or "").lower()
    action = str(activity.get("action") or "").lower()
    path = str(activity.get("path") or "").lower()
    target_name = str(activity.get("target_name") or "").lower()

    text = f"{module} {action} {path} {target_name}"

    if "report" in text or "báo cáo" in text or "bao cao" in text:
        return "Báo cáo"

    if "thu hồi" in text or "thu hoi" in text or "recover" in text or "recall" in text or "return" in text:
        return "Thu hồi"

    if "assign" in text or "cấp phát" in text or "cap phat" in text:
        return "Cấp phát"

    if "asset" in text or "tài sản" in text or "tai san" in text:
        return "Tài sản"

    if "user" in text or "người dùng" in text or "nguoi dung" in text:
        return "Người dùng"

    if "permission" in text or "phân quyền" in text or "phan quyen" in text:
        return "Phân quyền"

    if "mail" in text or "email" in text:
        return "Mail"

    if "activity" in text or "hoạt động" in text or "hoat dong" in text:
        return "Hoạt động"

    return "Hệ thống"


# Format thời gian cho file Excel xuất ra theo giờ Việt Nam
def format_activity_export_time(value):
    if not value:
        return ""

    if isinstance(value, datetime):
        local_value = convert_utc_to_vietnam_time(value)
        return local_value.strftime("%H:%M %d/%m/%Y")

    text = str(value).strip()

    if not text:
        return ""

    clean_text = text.replace("Z", "").replace("+00:00", "")

    formats = [
        "%Y-%m-%dT%H:%M:%S.%f",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%H:%M %d/%m/%Y",
        "%d/%m/%Y %H:%M",
    ]

    for fmt in formats:
        try:
            parsed = datetime.strptime(clean_text[:26], fmt)
            local_value = convert_utc_to_vietnam_time(parsed)
            return local_value.strftime("%H:%M %d/%m/%Y")
        except Exception:
            pass

    return text


# Tạo mô tả dễ đọc cho từng dòng trong file Excel
def build_activity_export_description(activity):
    action = activity.get("action") or ""
    target_name = activity.get("target_name") or ""
    path = activity.get("path") or ""

    if target_name:
        return f'{action} "{target_name}"'

    if path:
        return f"{action} ({path})"

    return action or "Không có mô tả"


# Xuất danh sách hoạt động ra file Excel
def get_activities_export(args, current_user_id):
    current_user = get_current_user_by_id(current_user_id)

    if not current_user:
        return None, None, {
            "success": False,
            "message": "Không tìm thấy user hiện tại"
        }, 404

    query, error_response, error_status = build_activity_query(args, current_user)

    if error_response:
        return None, None, error_response, error_status

    activities = (
        activities_collection
        .find(query)
        .sort("created_at", DESCENDING)
        .limit(10000)
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoạt động"

    headers = [
        "Người thực hiện",
        "Hành động",
        "Loại đối tượng",
        "Mô tả",
        "Thời gian",
    ]

    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True, color="000000")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="left", vertical="center")

    row_index = 2

    for activity in activities:
        full_name = (
            activity.get("full_name")
            or activity.get("email")
            or activity.get("employee_code")
            or "Chưa xác định"
        )

        action = activity.get("action") or "Thao tác hệ thống"
        activity_type = detect_activity_export_type(activity)
        description = build_activity_export_description(activity)
        created_at = (
            activity.get("created_at")
            or activity.get("updated_at")
            or activity.get("time")
        )

        ws.append([
            full_name,
            action,
            activity_type,
            description,
            format_activity_export_time(created_at),
        ])

        for cell in ws[row_index]:
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="left",
                vertical="top",
                wrap_text=True,
            )

        row_index += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 62
    ws.column_dimensions["E"].width = 20

    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 22

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename_time = datetime.now(APP_TIMEZONE).strftime("%Y%m%d_%H%M%S")
    filename = f"Activities_{filename_time}.xlsx"

    return output, filename, None, 200